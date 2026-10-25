"""Reorder the reviewed ONPG Worksheet without changing its scientific content.

The existing Word tables, figures, styles, and editable OMML equations are
moved unchanged. Only stage, experiment, question, and figure numbering are
remapped to match the six-stage simulation order.
"""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

PATH = Path(__file__).with_name("Worksheet.docx")
doc = Document(PATH)
body = doc.element.body
parts = list(body)
assert len(parts) == 193 and parts[-1].tag == qn("w:sectPr")

def words(i):
    return "".join(e.text or "" for e in parts[i].xpath(".//w:t"))

checks = {
  42:"Predictions Before Experiment 1",
  52:"Experiment 1: Compare",
  93:"Experiment 2: Free-Enzyme",
  110:"Experiment 2: Michaelis",
  121:"Experiment 2: Interpretation",
  131:"Experiment 3: Bead-Diameter",
  154:"Homework Problems",
  157:"Homework 1:",
  165:"Homework 2:",
}
for index, phrase in checks.items():
    assert phrase in words(index), (index, phrase)

stages = {1:4, 2:5, 3:6, 4:1, 5:2, 6:3}
experiments = {1:3, 2:1, 3:2}

def change_references(elem, swap_figure):
    for t in elem.xpath(".//w:t"):
        if not t.text:
            continue
        text = t.text
        def stage_replace(m):
            tail = ""
            if m.group(4):
                joiner = (" " + m.group(3) + " ") if m.group(3) == "and" else m.group(3)
                tail = joiner + str(stages[int(m.group(4))])
            return m.group(1) + " " + str(stages[int(m.group(2))]) + tail
        text = re.sub(r"\b(Stages?)\s+([1-6])(?:\s*([–-]|and)\s*([1-6]))?",stage_replace,text)
        text = re.sub(r"\bExperiment\s+([123])\b",
                      lambda m:"Experiment " + str(experiments[int(m.group(1))]),text)
        if swap_figure:
            text = re.sub(r"\bFig\.\s*([56])\b",
                          lambda m:"Fig. " + ("5" if m.group(1)=="6" else "6"),text)
        t.text = text

def renumber(index, old, new, homework=False):
    tokens = parts[index].xpath(".//w:t")
    first = next((t for t in tokens if (t.text or "").strip()),None)
    assert first is not None
    pattern = (r"^(\s*Question\s+)" + str(old) + r"(\.)" if homework
               else r"^(\s*)" + str(old) + r"(\.\s*)")
    assert re.search(pattern,first.text),(index,first.text)
    first.text = re.sub(pattern,lambda m:m.group(1)+str(new)+m.group(2),first.text,count=1)

for i,elem in enumerate(parts[:-1]):
    change_references(elem, 42<=i<131)

for number,index in enumerate([43,46,49,63,69,75,78,89],10):
    renumber(index,number-9,number)
for number,index in enumerate([105,111,116,122,128],1):
    renumber(index,number+8,number)
for number,index in enumerate([139,144,148,151],6):
    renumber(index,number+8,number)
for index,old,new in [(166,2,1),(176,3,2),(187,4,3),(158,1,4)]:
    renumber(index,old,new,True)

for index,old,new in [(165,2,1),(157,1,2)]:
    t=parts[index].xpath(".//w:t")[0]
    assert f"Homework {old}:" in t.text
    t.text=t.text.replace(f"Homework {old}:",f"Homework {new}:",1)

def new_page(elem):
    pPr=elem.find(qn("w:pPr"))
    if pPr is None:
        pPr=OxmlElement("w:pPr")
        elem.insert(0,pPr)
    if pPr.find(qn("w:pageBreakBefore")) is None:
        pPr.append(OxmlElement("w:pageBreakBefore"))

new_page(parts[42])   # bead predictions move after bead-design results
new_page(parts[154])  # new homework unit
new_page(parts[157])  # second homework assignment
# Prevent a near-empty page caused by the earlier second homework's break.
pPr=parts[165].find(qn("w:pPr"))
if pPr is not None:
    old_break=pPr.find(qn("w:pageBreakBefore"))
    if old_break is not None:
        pPr.remove(old_break)

# Keep all OOXML within each unit in its original order.
sequence=(parts[:42]+parts[92:131]+parts[131:154]+parts[42:92]
          +parts[154:157]+parts[165:192]+parts[157:165]+[parts[-1]])
assert len(sequence)==len(parts) and len({id(p) for p in sequence})==len(parts)
for elem in list(body):
    body.remove(elem)
for elem in sequence:
    body.append(elem)

doc.save(PATH)
review=Document(PATH)
alltext="\n".join(p.text for p in review.paragraphs)
assert "Experiment 1: Free-Enzyme Kinetics (Stage 1)" in alltext
assert "Experiment 2: Bead-Diameter Design (Stage 3)" in alltext
assert "Experiment 3: Compare the Two Bead Sizes (Stages 4–6)" in alltext
assert "Homework 1: ONPG Kinetics" in alltext
assert "Homework 2: ONP Measurement" in alltext
assert "Fig. 5. Initial-rate prediction graph" in alltext
assert "Fig. 6  Prediction axes" in alltext
assert len(review.tables)==5
assert len(review.element.body.xpath(".//m:oMath"))>=10
print("Worksheet staged order and editable equations verified.")
