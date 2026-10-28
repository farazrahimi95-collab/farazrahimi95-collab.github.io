"""Apply only professor-requested clarity edits to the reviewed ONPG worksheet.

Run AFTER worksheet reorder, homework formatting and seven-rate alignment.
Preserves all figures, equations, questions, scientific data and stage order.
Stage 3 stays UNCHANGED: professor explicitly requested discussion first.
"""
from pathlib import Path
from docx import Document

path = Path(__file__).with_name("Worksheet.docx")
doc = Document(path)
math_before = len(doc.element.body.xpath(".//m:oMath"))
table_before = len(doc.tables)
paragraph_before = len(doc.paragraphs)

def change_exact(start, replacement):
    hits = [p for p in doc.paragraphs if p.text.startswith(start)]
    if len(hits) != 1:
        raise ValueError(f"Expected one paragraph starting {start!r}; found {len(hits)}")
    p = hits[0]
    if len(p.runs) != 1:
        raise ValueError(f"Cannot preserve formatting of {start!r}; refusing edit")
    p.runs[0].text = replacement

# Simplify the blanking explanation without altering the calibration.
change_exact("Before each run, the experimentalist obtains a baseline absorbance",
    "Before ONP measurement, a clear ONPG reference is used to zero the spectrophotometer and remove background absorbance. The corrected reading at 415 nm is used to calculate ONP concentration.")
change_exact("Before measuring ONP, blank the spectrophotometer using a clear liquid reference.",
    "The reference cuvette shown on the left contains clear liquid; it is not empty. The seven Stage 1 assay cuvettes contain equal enzyme amounts but different initial ONPG concentrations. Their yellow color represents ONP formed during the scan.")

# Clarify the experiment actions in the original four-step procedure.
change_exact("1. Open Stage 4 (2.8 mm beads). Click",
    "1. In Stage 4 (2.8 mm beads), click “Set autosampler · six sample times.” This records the initial reading at 0 min and schedules five additional readings.")
change_exact("2. Click “Start 340-min run (1-minute playback).”",
    "2. Click “Start 340-min run (1-minute playback).” The simulation displays 340 equivalent laboratory minutes in 60 seconds. Observe the smooth ONP concentration–time curve while the table records readings at 0, 28, 85, 170, 255, and 340 min.")
change_exact("3. Continue to Stage 5 (4.6 mm beads).",
    "3. In Stage 5 (4.6 mm beads), repeat the same autosampler setup and 60-second run.")
change_exact("4. Open Stage 6 to compare the two saved runs.",
    "4. Open Stage 6 to compare the six readings and complete curves from both saved runs. No new reactor run is performed in Stage 6.")

assert len(doc.paragraphs) == paragraph_before
assert len(doc.tables) == table_before == 5
assert len(doc.element.body.xpath(".//m:oMath")) == math_before
s = "\n".join(p.text for p in doc.paragraphs)
for t in ("Experiment 1: Free-Enzyme Kinetics (Stage 1)",
          "Experiment 1: Michaelis–Menten Analysis (Stage 2)",
          "Experiment 2: Bead-Diameter Design (Stage 3)",
          "Experiment 3: Compare the Two Bead Sizes (Stages 4–6)"):
    assert t in s
assert s.index("Experiment 1: Free-Enzyme Kinetics (Stage 1)") < s.index("Experiment 1: Michaelis–Menten Analysis (Stage 2)") < s.index("Experiment 2: Bead-Diameter Design (Stage 3)") < s.index("Experiment 3: Compare the Two Bead Sizes (Stages 4–6)")
assert "15. Which bead has the greater percentage decrease" in s
doc.save(path)
print("Professor-requested Stage 4/5 procedure and blanking text clarified; Stage 3 unchanged.")
