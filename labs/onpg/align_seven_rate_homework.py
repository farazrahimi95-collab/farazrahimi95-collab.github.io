"""Align Stage 2 reciprocal-plot homework with all seven Stage 1 assay rates.
This follows reorder_worksheet.py and format_homework_questions.py.
Preserve native Word equation objects and existing run-level subscripts.
"""
from pathlib import Path
from docx import Document

path = Path(__file__).with_name("Worksheet.docx")
doc = Document(path)
math_count = len(doc.element.body.xpath(".//m:oMath"))
table_count = len(doc.tables)

def find_one(prefix):
    matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    assert len(matches) == 1, (prefix, len(matches))
    return matches[0]

def update_run(p, old, new):
    hits = [r for r in p.runs if old in r.text]
    assert len(hits) == 1, (old, len(hits))
    hits[0].text = hits[0].text.replace(old, new)

intro = find_one("Complete these problems after using the simulation.")
update_run(intro,
    "Use the six supplied laboratory rates for regression; do not treat the additional 39.8 mM model-based assay as an independent measurement.",
    "Use all seven Stage 1 initial rates in the kinetics and reciprocal-plot questions.")

question = find_one("Question 2. Using only the six supplied laboratory initial rates")
update_run(question,
    "Using only the six supplied laboratory initial rates (0.622–19.9 mM)",
    "Using all seven Stage 1 initial rates (0.622–39.8 mM)")
update_run(question,
    " from the fitted slope and intercept. The additional 39.8 mM virtual rate is excluded from this experimental fit.",
    " from the fitted slope and intercept.")

figure = find_one("Fig. 7. Reciprocal plot for the six supplied laboratory assays.")
update_run(figure,
    "Fig. 7. Reciprocal plot for the six supplied laboratory assays. Exclude the simulated 39.8 mM point.",
    "Fig. 7. Reciprocal plot for all seven Stage 1 initial rates.")

assert len(doc.tables) == table_count == 5
assert len(doc.element.body.xpath(".//m:oMath")) == math_count
doc.save(path)

checked = Document(path)
alltext = "\n".join(p.text for p in checked.paragraphs)
assert "Using all seven Stage 1 initial rates (0.622–39.8 mM)" in alltext
assert "Fig. 7. Reciprocal plot for all seven Stage 1 initial rates." in alltext
assert "exclude" not in "\n".join(p.text for p in checked.paragraphs if p.text.startswith(("Question 2.", "Fig. 7."))).lower()
assert len(checked.element.body.xpath(".//m:oMath")) == math_count
print(f"Seven-rate homework aligned; preserved {math_count} editable equations and {table_count} tables.")
