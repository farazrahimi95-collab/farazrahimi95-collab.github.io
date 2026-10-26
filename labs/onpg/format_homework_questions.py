"""Format ONPG worksheet homework as one section with four numbered questions.
Run after reorder_worksheet.py, before PDF conversion.
"""
from pathlib import Path
from docx import Document

path = Path(__file__).with_name("Worksheet.docx")
doc = Document(path)
headings = {}
questions = []
for paragraph in doc.paragraphs:
    text = paragraph.text.strip()
    if text == "Homework Problems":
        headings["main"] = paragraph
    elif text.startswith("Homework 1:"):
        headings["first"] = paragraph
    elif text.startswith("Homework 2:"):
        headings["second"] = paragraph
    elif text.startswith("Question "):
        questions.append(paragraph)

assert len(headings) == 3, "Worksheet homework structure changed"
assert [p.text.split(".",1)[0] for p in questions] == [
    "Question 1", "Question 2", "Question 3", "Question 4"
], "Questions must be numbered 1–4 in sequence"
before_math = len(doc.element.body.xpath(".//m:oMath"))
before_tables = len(doc.tables)

headings["main"].runs[0].text = "Homework Questions"
questions[3].paragraph_format.page_break_before = True
for question in questions:
    question.style = doc.styles["Homework Question"]
for name in ("first", "second"):
    element = headings[name]._element
    element.getparent().remove(element)

assert len(doc.tables) == before_tables
assert len(doc.element.body.xpath(".//m:oMath")) == before_math

doc.save(path)
check = Document(path)
lines = [p.text.strip() for p in check.paragraphs]
assert lines.count("Homework Questions") == 1
assert not any(p.startswith(("Homework 1:", "Homework 2:", "Homework Problems")) for p in lines)
assert len(check.element.body.xpath(".//m:oMath")) == before_math
print("Homework Questions formatted: Question 1–4, original science preserved.")
