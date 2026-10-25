from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image,ImageDraw,ImageFont
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree

ROOT=Path(__file__).resolve().parent
SRC=ROOT/"ONPG_Virtual_Lab_Worksheet.docx"
OUT=ROOT/"Worksheet.docx"
D=Document(SRC)

def paragraph(fragment):
    for p in D.paragraphs:
        if fragment.lower() in p.text.lower():return p
    return None

def change(fragment,new,required=False):
    p=paragraph(fragment)
    if not p:
        if required:raise RuntimeError("Missing expected worksheet text: "+fragment)
        print("optional paragraph missing:",fragment)
        return
    p.text=new

def delete(elem):
    e=elem._element
    if e.getparent() is not None:e.getparent().remove(e)

def add_before(anchor,text,style="Normal"):
    p=anchor.insert_paragraph_before(style=style);p.text=text
    return p

def math_r(txt):
    r=OxmlElement("m:r");t=OxmlElement("m:t");t.text=str(txt);r.append(t);return r

def omml(anchor,chunks):
    p=anchor.insert_paragraph_before(style="Normal")
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(6)
    para=OxmlElement("m:oMathPara");m=OxmlElement("m:oMath")
    for c in chunks:
        if isinstance(c,tuple):
            frac=OxmlElement("m:f")
            for name,value in [("m:num",c[0]),("m:den",c[1])]:
                el=OxmlElement(name);el.append(math_r(value));frac.append(el)
            m.append(frac)
        else:m.append(math_r(c))
    para.append(m);p._p.append(para)

def fnt(bold=False,size=21):
    try:return ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans"+("-Bold" if bold else "")+".ttf",size)
    except:return ImageFont.load_default()

# Preserve the existing WSU worksheet and all its source figures, replacing only
# passages invalidated by the current seven-assay simulation.
change("Development date:","Worksheet revised October 2026")
change("Use six initial rates","Use the six supplied initial rates and one 39.8 mM simulated assay to interpret Michaelis-Menten behavior.")
change("Before measuring ONP, the spectrophotometer is blanked","A clear liquid reference is used to blank the spectrophotometer before measuring ONP. Stage 4 uses seven assay cuvettes with equal enzyme amounts and different initial ONPG concentrations.")
change("Fig. 3  Clear reference blank","Fig. 3  Clear liquid reference and seven Stage 4 assay cuvettes. Yellow coloring represents ONP formation.")
change("Set autosampler · six worksheet times","1. Open Stage 1 (2.8 mm beads) and select Set autosampler · six sample times.")
change("2-minute playback","2. Click Start 340-min run (1-minute playback). Six automatic readings are taken at 0, 28, 85, 170, 255 and 340 equivalent laboratory minutes; observe the full continuous conversion.")
change("3. Continue to Stage 2","3. Continue to Stage 2 (4.6 mm beads); collect the same six readings in its 60-second playback.")
change("The internal effectiveness factor","The internal effectiveness factor, ηᵢ, is the whole-bead reaction rate divided by that expected if the bead were at its surface substrate concentration throughout. Smaller values indicate stronger diffusion limitations.")
change("Purpose: determine the initial rate at six","Purpose: compare seven ONPG initial concentrations at equal enzyme amount, volume, pH and temperature. The first six rates are supplied laboratory data; 39.8 mM is a simulated rate.")
change("The six slopes become","Seven initial-rate slopes appear in the Stage 5 Michaelis-Menten analysis.")
change("Load six cuvettes","1. Click Load seven cuvettes in Stage 4.")
change("All six ONP-concentration-versus-time lines","3. Click Start 4.5-s scan. Seven ONP versus time lines appear on the same graph.")
change("record the six slopes","4. Click Show initial rates and record all seven rates in the table.")
change("Compare the slopes at 0.622","9. Compare the rates at 0.622 versus 4.97 mM, then at 19.9 versus 39.8 mM ONPG. Why does the rate increase more slowly as ONPG becomes abundant?")
change("Before opening the Stage 5 fitted","10. Predict the shape of the initial-rate curve from 0 to 40 mM ONPG. Sketch the curve on Fig. 6 and explain the expected saturation.")
change("Fig. 6  Prediction","Fig. 6  Prediction and initial-rate graph from 0 to 40 mM. Plot all seven Stage 4 rates.")
change("Plot your six actual","11. Plot all seven Stage 4 rates on Fig. 6 and compare with the Stage 5 curve.")
change("Using Figs. 6 and 6b","12. Using Fig. 6 and Stage 5, explain how initial rate changes with initial ONPG concentration.")
change("At low  values","At low ONPG concentrations, why is initial rate approximately proportional to ONPG?")
change("When  >>","At high ONPG concentrations, why does adding substrate produce progressively smaller rate increases?")
change("Using Figs. 6","12. Describe how initial ONPG concentration affects initial rate, particularly near saturation.")
change("for example, 170 min","for example, 170 min")
change("0 and 28.3","0 and 28")
change("Question 3. Use the Michaelis–Menten","Question 3. Plot 1/v₀ against 1/C_ONPG,0 using only the six supplied laboratory rates (0.622-19.9 mM). Use the fitted slope and intercept to determine Vₘₐₓ and Kₘ. Exclude the simulated 39.8 mM value from this experimental regression.")
change("Fig. 7  Blank reciprocal","Fig. 7  Reciprocal graph for six supplied ONPG initial-rate measurements.")
change("Compare the parameters calculated from your reciprocal plot","The six-point reciprocal regression need not match the Stage 5 anchored reference (Kₘ = 2.00 mM; Vₘₐₓ = 0.002510 mM/s). The two methods use different assumptions and error weighting.")
change("Question 4. Over the first seconds","Question 4. Compare the rates at 9.95, 19.9 and 39.8 mM. The 39.8 mM result is model-generated. Explain why the change becomes smaller near saturation.")
change("Complete these problems after the in-class simulation","Complete these problems after running the simulation. Show equations, numerical substitutions and units. Use only the six supplied laboratory rates for experimental regression; do not count the 39.8 mM simulated assay as another measurement.")
# Chart-independent tables.
for table in D.tables:
    if len(table.columns)==3 and len(table.rows)==7 and table.cell(2,0).text.strip().startswith("28."):
        table.cell(2,0).text="28"
    if len(table.columns)==5 and any("Large laboratory bead" in c.text for row in table.rows for c in row.cells):
        for row in table.rows:
            if row.cells[0].text.startswith("Large laboratory"):row.cells[2].text="0.2885"
            if row.cells[0].text.startswith("Small laboratory"):row.cells[2].text="0.4890"
    if len(table.columns)==2 and table.cell(0,0).text.strip().startswith("Initial") and len(table.rows) in (7,8):
        vals=[row.cells[0].text for row in table.rows]
        if "19.9" in vals and "39.8" not in vals:
            row=table.add_row();row.cells[0].text="39.8";row.cells[1].text=""
        table.cell(0,0).text="Initial ONPG, C_ONPG,0 [mM]"
        table.cell(0,1).text="Initial rate, v₀ [mM/s]"
# Remove the historical three-value extrapolation block, including its table and second chart.
start=paragraph("Stage 5: High-Concentration Model Extension")
stop=paragraph("Experiment 2: Interpretation of Kinetics")
if start and stop:
    body=list(D.element.body)
    first=body.index(start._element);last=body.index(stop._element)
    for obj in body[first:last]:obj.getparent().remove(obj)
    add_before(stop,"The 39.8 mM assay is generated from the Michaelis-Menten reference (Kₘ = 2.00 mM, Vₘₐₓ = 0.002510 mM/s). It is anchored to the supplied 19.9 mM value and assumes no substrate inhibition.",style="Small note")
    omml(stop,["v₀ = ",("Vₘₐₓ C_ONPG,0","Kₘ + C_ONPG,0")])
# Remove the old eight-step homework derivation and build a concise native-editable math task.
start=paragraph("Student Review: Deriving the Michaelis")
stop=paragraph("Question 3.")
if start and stop:
    body=list(D.element.body);first=body.index(start._element);last=body.index(stop._element)
    for obj in body[first:last]:obj.getparent().remove(obj)
    add_before(stop,"Question 2. In the free β-galactosidase assay, ONPG is the substrate and ONP is the measured product. Use the relationship below; a complete enzyme-mechanism derivation is not required.",style="Homework Question")
    omml(stop,["v₀ = ",("Vₘₐₓ C_ONPG,0","Kₘ + C_ONPG,0")])
    add_before(stop,"a. Define Vₘₐₓ and Kₘ physically and state their units.")
    add_before(stop,"b. Take the reciprocal and identify slope m and intercept b in the graph of 1/v₀ versus 1/C_ONPG,0.")
    omml(stop,[("1","v₀")," = ",("Kₘ","Vₘₐₓ")," × ",("1","C_ONPG,0")," + ",("1","Vₘₐₓ")])
    add_before(stop,"c. Express Vₘₐₓ and Kₘ in terms of the reciprocal-plot slope and intercept.")
    omml(stop,["Vₘₐₓ = ",("1","b"),"     Kₘ = ",("m","b")])
    add_before(stop,"Calculations and interpretation:")
    add_before(stop,"")
    add_before(stop,"")
# Figure 3 contains seven cuvettes, not a reused six-cuvette graphic.
figcap=paragraph("Fig. 3")
if figcap:
    prev=figcap._element.getprevious()
    if prev is not None:
        from docx.text.paragraph import Paragraph
        figpara=Paragraph(prev,figcap._parent)
        if figpara._p.xpath(".//a:blip"):
            figpara.clear()
            pic=ROOT/"worksheet_figure3.png"
            img=Image.new("RGB",(1300,270),"#eff7f6");dr=ImageDraw.Draw(img)
            labels=["Blank","0.622","1.24","2.49","4.97","9.95","19.9","39.8"]
            fills=["#fff","#fff3ca","#ffe9a6","#ffe08b","#fcd56b","#ffcc52","#eeb844","#e6aa36"]
            for i,(label,fill) in enumerate(zip(labels,fills)):
                x=105+158*i
                dr.rounded_rectangle((x-31,54,x+31,178),radius=8,fill="#eaf2f2",outline="#78a8a8",width=4)
                dr.rounded_rectangle((x-25,95,x+25,171),radius=4,fill=fill)
                dr.text((x-31,195),label,font=ImageFont.load_default(),fill="#164253")
                dr.text((x-16,221),"liquid" if i==0 else "mM",font=ImageFont.load_default(),fill="#2c5360")
            img.save(pic)
            figpara.add_run().add_picture(str(pic),width=Inches(6.3))
# Save and patch editable Excel-style chart axes plus legacy math 28.3 literals.
D.save(OUT)
with ZipFile(OUT,"r") as z:files={p:z.read(p) for p in z.namelist()}
chart=files.get("word/charts/chart3.xml")
if chart:
    root=etree.fromstring(chart); ns="{http://schemas.openxmlformats.org/drawingml/2006/chart}"
    for ax in root.findall(".//"+ns+"valAx"):
        scaling=ax.find(ns+"scaling")
        lim=scaling.find(ns+"max") if scaling is not None else None
        if lim is not None and lim.get("val")=="20":
            lim.set("val","40");u=ax.find(ns+"majorUnit")
            if u is not None:u.set("val","10")
    files["word/charts/chart3.xml"]=etree.tostring(root,encoding="UTF-8",xml_declaration=True)
files["word/document.xml"]=files["word/document.xml"].replace(b"28.33",b"28").replace(b"28.3",b"28")
with ZipFile(OUT,"w",ZIP_DEFLATED) as z:
    for name,blob in files.items():z.writestr(name,blob)
assert OUT.stat().st_size>100000
print("Generated",OUT.name,OUT.stat().st_size,"bytes; native editable Word mathematics preserved")
