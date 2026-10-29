"""Arrange reviewed ONPG worksheet content by the actual six simulation stages.

Run after reorder_worksheet.py, format_homework_questions.py,
align_seven_rate_homework.py and apply_professor_clarifications.py.
Move original XML blocks intact, preserving figures, tables, math and questions.
Only stage headings, stage instructions and navigation text are adjusted.
"""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

path = Path(__file__).with_name("Worksheet.docx")
doc = Document(path)
body = doc.element.body
old = list(body)
assert old[-1].tag == qn("w:sectPr")
table_count = len(doc.tables)
math_count = len(body.xpath(".//m:oMath"))
assert table_count == 5 and math_count >= 10

def words(e):
    return "".join(t.text or "" for t in e.xpath(".//w:t"))

def locate(prefix):
    hits = [i for i,e in enumerate(old) if e.tag == qn("w:p") and words(e).startswith(prefix)]
    assert len(hits)==1,(prefix,hits)
    return hits[0]

def para(prefix):
    return Paragraph(old[locate(prefix)],doc)

h = {
 "s1":locate("Experiment 1: Free-Enzyme Kinetics (Stage 1)"),
 "s2":locate("Experiment 1: Michaelis–Menten Analysis (Stage 2)"),
 "s3":locate("Experiment 2: Bead-Diameter Design (Stage 3)"),
 "predict":locate("Predictions Before Experiment 3"),
 "group":locate("Experiment 3: Compare the Two Bead Sizes (Stages 4–6)"),
 "paired":locate("Experiment 3: Paired Results and Performance"),
 "comparison":locate("Bead-Size Comparison and Internal Mass Transfer"),
 "internal":locate("Experiment 3: Internal Concentration and Effectiveness"),
 "homework":locate("Homework Questions"),
}
assert list(h.values()) == sorted(h.values())

def fresh(text,style="Normal",bold_prefix=None,keep=False):
    p=doc.add_paragraph(style=style)
    if bold_prefix is not None:
        assert text.startswith(bold_prefix)
        p.add_run(bold_prefix).bold=True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    p.paragraph_format.keep_with_next=keep
    body.remove(p._element)
    return p._element

def heading(t):
    return fresh(t,style="Heading 2",keep=True)

def next_stage(n):
    return fresh(
        f"When this stage is complete, click “Continue to Stage {n}” at the bottom of the simulation.",
        bold_prefix="When this stage is complete,")

# Stage 2 procedure before Questions 2–5, matching the actual simulation UI.
s2procedure=[
    heading("Procedure"),
    fresh("1. After completing Stage 1, click “Continue to Stage 2” at the bottom of the simulation."),
    fresh("2. Before building the plot, complete the prediction in Question 2 on Fig. 5. Then click “Build kinetics plot” in the simulation."),
    fresh("3. Click “Show half-maximum guide” in the simulation and use the displayed curve to answer Questions 3–5."),
]

# Stage 3 procedure remains scientifically identical; name the real Save button.
for r in para("1. In Stage 3, use the slider or exact-diameter field").runs:
    if "then save the result" in r.text:
        r.text=r.text.replace("then save the result","then click “Save result”")
for r in para("2. Set the diameter to 0.3 mm").runs:
    if "calculate, and save its result" in r.text:
        r.text=r.text.replace("calculate, and save its result",
                              "click “Calculate” and then “Save result”")

# Existing combined 4-step procedure is split across Stages 4, 5, and 6.
def one_of(*prefixes):
    for p in prefixes:
        hits=[i for i,e in enumerate(old) if e.tag==qn("w:p") and words(e).startswith(p)]
        if len(hits)==1:
            return Paragraph(old[hits[0]],doc)
    raise AssertionError(("Missing procedure",prefixes))

run4a=one_of("1. In Stage 4 (2.8 mm beads)","1. Open Stage 4 (2.8 mm beads)")
run4b=para("2. Click “Start 340-min run (1-minute playback).”")
run5=one_of("3. In Stage 5 (4.6 mm beads)","3. Continue to Stage 5 (4.6 mm beads)")
run6=para("4. Open Stage 6 to compare")

run4a.text=(
    "1. After Stage 3, click “Continue to Stage 4.” In Stage 4 (2.8 mm beads), "
    "click “Set autosampler · six sample times” to capture the zero baseline and five later readings."
)
run4b.text="3. "+run4b.text[3:]
run5.text=(
    "1. Click “Continue to Stage 5” at the bottom of Stage 4. In Stage 5 (4.6 mm beads), "
    "click “Set autosampler · six sample times,” then “Start 340-min run (1-minute playback).” "
    "The same six equivalent times are recorded for comparison with Stage 4."
)
run6.text=(
    "1. Click “Continue to Stage 6” at the bottom of Stage 5. Compare the six readings and "
    "complete ONP concentration–time curves from both runs. No new reactor run is performed in Stage 6."
)

pred=para("Predictions Before Experiment 3")
pred.text="Stage 4 Questions: Predictions Before the Run"
pred.style=doc.styles["Heading 2"]
pred.paragraph_format.page_break_before=False

paired=para("Experiment 3: Paired Results and Performance")
paired.text="Stage 5 Questions: Paired Results and Performance"
paired.style=doc.styles["Heading 2"]
paired.paragraph_format.page_break_before=True

compare=para("Bead-Size Comparison and Internal Mass Transfer")
compare.text="Stage 6: Bead-Size Comparison and Internal Mass Transfer"
compare.style=doc.styles["Heading 2"]

inner=para("Experiment 3: Internal Concentration and Effectiveness")
inner.text="Stage 6 Questions: Internal Concentration and Effectiveness"
inner.style=doc.styles["Heading 2"]
inner.paragraph_format.page_break_before=False

interp=para("Experiment 1: Interpretation of Kinetics")
interp.text="Stage 2 Questions: Interpretation of Kinetics"
interp.style=doc.styles["Heading 2"]

# Original section XML is transferred, never recreated (equations and figures
# therefore remain editable and original question text stays untouched).
before=old[:h["s1"]]
stage1=old[h["s1"]:h["s2"]]
stage2=old[h["s2"]:h["s3"]]
stage3=old[h["s3"]:h["predict"]]
predictions=old[h["predict"]:h["group"]]
group=old[h["group"]:h["paired"]]
pairedsection=old[h["paired"]:h["comparison"]]
comparison=old[h["comparison"]:h["internal"]]
internal=old[h["internal"]:h["homework"]]
homework=old[h["homework"]:-1]

i4=group.index(run4a._element)
i4b=group.index(run4b._element)
i5=group.index(run5._element)
i6=group.index(run6._element)
assert i4<i4b<i5<i6
intro=group[:i4]
steps4=group[i4:i5]
steps4.insert(steps4.index(run4b._element),
              fresh("2. Before starting the Stage 4 run, answer Questions 10–12 and sketch your prediction on Fig. 6 below."))
steps5=group[i5:i6]
steps6=group[i6:]

ordered = (
    before + stage1 + [next_stage(2)] +
    stage2[:1] + s2procedure + stage2[1:] + [next_stage(3)] +
    stage3 + [next_stage(4)] +
    intro[:-1] + [heading("Stage 4: Small-Bead Reactor Run (2.8 mm)")] +
    intro[-1:] + steps4 + predictions + [next_stage(5)] +
    [heading("Stage 5: Large-Bead Reactor Run (4.6 mm)"),
     heading("Procedure")] + steps5 + pairedsection + [next_stage(6)] +
    comparison[:1] + [heading("Procedure")] + steps6 +
    comparison[1:] + internal + homework + [old[-1]]
)
assert len(ordered)==len(set(map(id,ordered))), "Duplicated worksheet content"
assert set(map(id,old)).issubset(set(map(id,ordered))), "Lost worksheet content"

for e in list(body):body.remove(e)
for e in ordered:body.append(e)

# Keep Question 11 with its prediction figure; start the long Q16+Q17 block
# on the following page to prevent the previous orphaned Question 17.
para("11. Before running the simulation, sketch").paragraph_format.keep_with_next=True
para("16. What is happening to ONPG inside the larger bead").paragraph_format.page_break_before=True

t="\n".join(p.text for p in doc.paragraphs)
heads=[
 "Experiment 1: Free-Enzyme Kinetics (Stage 1)",
 "Experiment 1: Michaelis–Menten Analysis (Stage 2)",
 "Experiment 2: Bead-Diameter Design (Stage 3)",
 "Stage 4: Small-Bead Reactor Run (2.8 mm)",
 "Stage 5: Large-Bead Reactor Run (4.6 mm)",
 "Stage 6: Bead-Size Comparison and Internal Mass Transfer",
 "Homework Questions",
]
positions=[t.index(s) for s in heads]
assert positions==sorted(positions)
questions=[
 "1. Compare the slopes","2. Predict the shape","3. Plot the seven",
 "4. Using Fig. 5","5. Calculate V","6. For the assigned diameter",
 "7. Calculate the ratio","8. Would there","9. Why should",
 "10. Before running the simulation","11. Before running the simulation",
 "12. What are the reasons","13. At the same time",
 "14. Which bead preparation","15. Which bead has the greater percentage decrease",
 "16. What is happening to ONPG","17. Return to your prediction",
]
stages=[0,1,1,1,1,2,2,2,2,3,3,3,4,5,5,5,5]
for i,(q,st) in enumerate(zip(questions,stages),1):
    matches=[p for p in doc.paragraphs if p.text.startswith(q)]
    assert len(matches)==1,(i,q,len(matches))
    where=t.index(matches[0].text)
    assert positions[st]<where<positions[st+1],("Wrong-stage question",i)
assert len(doc.tables)==table_count
assert len(body.xpath(".//m:oMath"))==math_count
assert t.count("Homework Questions")==1
for n in range(2,7):
    assert f"Continue to Stage {n}" in t

doc.save(path)
print(f"Six-stage worksheet verified: 17 questions, {table_count} tables, {math_count} editable math expressions.")
