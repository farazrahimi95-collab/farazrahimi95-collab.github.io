"""Final approved worksheet revisions, after organize_worksheet_by_stage.py.
Remove classroom Question 3 and homework Question 4; renumber classroom questions.
Restore one Stages 4–6 procedure, then prediction, then comparison questions.
Preserve all remaining scientific content, drawings, tables and equations.
"""
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from pathlib import Path
import re

path=Path(__file__).with_name("Worksheet.docx")
doc=Document(path)
body=doc.element.body
parts=list(body)
assert parts[-1].tag==qn("w:sectPr") and len(doc.tables)==5
math_before=len(body.xpath(".//m:oMath"))
assert math_before==18

def words(e):
    return "".join(t.text or "" for t in e.xpath(".//w:t"))
def pos(prefix):
    hits=[i for i,e in enumerate(parts) if e.tag==qn("w:p") and words(e).startswith(prefix)]
    assert len(hits)==1,(prefix,hits)
    return hits[0]
def el(prefix): return parts[pos(prefix)]
def par(prefix): return Paragraph(el(prefix),doc)

# Q3's answer space and its associated explanatory note are removed.
# Preserve the following editable Michaelis–Menten equation.
start=pos("3. Plot the seven Stage 1 rates")
finish=pos("Stage 2 Questions: Interpretation of Kinetics")
assert len(parts[finish-1].xpath(".//m:oMath"))==1
removed=set(parts[start:finish-1])
caption=par("Fig. 5. Initial-rate prediction graph")
assert "then plot the seven Stage 1 rates." in caption.text
for r in caption.runs:
    r.text=r.text.replace(", then plot the seven Stage 1 rates.", ".")
step=par("3. Click “Show half-maximum guide”")
assert "Questions 3–5" in step.text
for r in step.runs:
    r.text=r.text.replace("Questions 3–5","Questions 3–4")

qs=["4. Using Fig. 5 and Stage 2","5. Calculate V",
    "6. For the assigned diameter","7. Calculate the ratio","8. Would there",
    "9. Why should we immobilize","10. Before running the simulation, which bead",
    "11. Before running the simulation, sketch","12. What are the reasons",
    "13. At the same time points","14. Which bead preparation",
    "15. Which bead has the greater percentage decrease",
    "16. What is happening to ONPG","17. Return to your prediction"]
for old,prefix in enumerate(qs,4):
    p=par(prefix)
    for r in p.runs:
        if re.match(r"^\s*"+str(old)+r"\.",r.text):
            r.text=re.sub(r"^([ \t]*)"+str(old)+r"(?=\.)",lambda m:m.group(1)+str(old-1),r.text,count=1)
            break
    else:
        raise AssertionError(("Missing numeric run",old,prefix))
    assert p.text.startswith(str(old-1)+". ")

hw=pos("Question 4. Verify the absorbance calibration")
assert pos("Question 3. Compare the rates")<hw
removed.update(parts[hw:-1])

begin=pos("Experiment 3: Compare the Two Bead Sizes (Stages 4–6)")
end=pos("Homework Questions")
sec=parts[begin:end]
group=el("Experiment 3: Compare the Two Bead Sizes (Stages 4–6)")
purpose=el("Purpose: Compare")
procedure=parts[pos("Stage 4: Small-Bead Reactor Run (2.8 mm)")+1]
assert words(procedure)=="Procedure"
a=el("1. After Stage 3, click “Continue to Stage 4.”")
b=el("2. Before starting the Stage 4 run, answer Questions 10–12")
c=el("3. Click “Start 340-min run (1-minute playback).”")
d=el("1. Click “Continue to Stage 5” at the bottom of Stage 4.")
e=el("1. Click “Continue to Stage 6” at the bottom of Stage 5.")
Paragraph(a,doc).text=(
    "1. In Stage 4 (2.8 mm beads), click “Set autosampler · six sample times.” "
    "This captures the initial reading at 0 min and schedules five more readings.")
p=Paragraph(c,doc)
p.text="2. "+p.text.split(". ",1)[1]
Paragraph(d,doc).text=(
    "3. Click “Continue to Stage 5” in the simulation. For the 4.6 mm beads, click "
    "“Set autosampler · six sample times,” then “Start 340-min run (1-minute playback).” "
    "The same six equivalent laboratory times are recorded.")
Paragraph(e,doc).text=(
    "4. Click “Continue to Stage 6” in the simulation to compare the six readings "
    "and complete ONP concentration–time curves for both bead sizes. "
    "No additional reactor run is performed in Stage 6.")
Paragraph(b,doc).text=(
    "After reading the complete procedure, answer prediction Questions 9–11 below "
    "before clicking “Start 340-min run” in Stage 4. Then perform Stages 4–6 "
    "and answer the remaining comparison questions.")
for old,new in [
    ("Stage 4 Questions: Predictions Before the Run","Predictions Before the Reactor Runs"),
    ("Stage 5 Questions: Paired Results and Performance","Paired Results and Performance"),
    ("Stage 6: Bead-Size Comparison and Internal Mass Transfer","Bead-Size Comparison and Internal Mass Transfer"),
    ("Stage 6 Questions: Internal Concentration and Effectiveness","Internal Concentration and Effectiveness"),
]:
    p=par(old)
    p.text=new
    p.paragraph_format.page_break_before=False

pred=el("Predictions Before the Reactor Runs")
paired=el("Paired Results and Performance")
comparison=el("Bead-Size Comparison and Internal Mass Transfer")
internal=el("Internal Concentration and Effectiveness")
nav5=el("When this stage is complete, click “Continue to Stage 5”")
nav6=el("When this stage is complete, click “Continue to Stage 6”")
combined=(
    [group,purpose,procedure,a,c,d,e,b]
    +sec[sec.index(pred):sec.index(nav5)]
    +sec[sec.index(paired):sec.index(nav6)]
    +[comparison]
    +sec[sec.index(e)+1:sec.index(internal)]
    +sec[sec.index(internal):]
)
assert len(combined)==len(set(map(id,combined)))
# Discard the obsolete page break that once split Stage 6 into two pages.
for item in combined:
    if item.tag==qn("w:p") and not words(item).strip() and item.xpath('.//w:br[@w:type="page"]'):
        removed.add(item)
new=[x for x in parts[:begin] if x not in removed]
new.extend(x for x in combined if x not in removed)
new.extend(x for x in parts[end:-1] if x not in removed)
new.append(parts[-1])
assert len(new)==len(set(map(id,new)))
for x in list(body):body.remove(x)
for x in new:body.append(x)
par("15. What is happening to ONPG").paragraph_format.page_break_before=False

text="\n".join(p.text for p in doc.paragraphs)
checks=["1. Compare the slopes","2. Predict the shape","3. Using Fig. 5",
        "4. Calculate V","5. For the assigned diameter","6. Calculate the ratio",
        "7. Would there","8. Why should we immobilize",
        "9. Before running the simulation, which bead",
        "10. Before running the simulation, sketch","11. What are the reasons",
        "12. At the same time points","13. Which bead preparation",
        "14. Which bead has the greater percentage decrease",
        "15. What is happening to ONPG","16. Return to your prediction",
        "Question 1.","Question 2.","Question 3."]
indices=[text.index(t) for t in checks]
assert indices==sorted(indices)
assert "Question 4." not in text and "3. Plot the seven Stage 1 rates" not in text
assert "then plot the seven Stage 1 rates" not in text and "Questions 3–5" not in text
assert len(doc.tables)==5 and len(body.xpath(".//m:oMath"))==math_before-3
assert text.index('4. Click “Continue to Stage 6”')<text.index("Predictions Before the Reactor Runs")
assert text.index("Predictions Before the Reactor Runs")<text.index("Paired Results and Performance")<text.index("Bead-Size Comparison and Internal Mass Transfer")
doc.save(path)
print("Final ONPG worksheet: 16 classroom questions, 3 homework questions, 5 tables and 15 editable equations.")
