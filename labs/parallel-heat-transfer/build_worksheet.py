from pathlib import Path
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement

BASE = Path(__file__).resolve().parent
DOCX = BASE / "Parallel_Heat_Transfer_Worksheet.docx"

doc = Document(DOCX)

def norm(text):
    return " ".join((text or "").replace("\u201c", '"').replace("\u201d", '"').replace("\u2013", "-").replace("\u2014", "-").split()).lower()

def find_para(*needles):
    n = [x.lower() for x in needles]
    for p in doc.paragraphs:
        t = norm(p.text)
        if all(x in t for x in n):
            return p
    raise RuntimeError("Paragraph not found: " + " | ".join(needles))

def replace_para(p, text):
    p.text = text

def append_once(p, text, marker):
    if marker.lower() not in norm(p.text):
        p.add_run(text)

def remove_para(p):
    el = p._element
    el.getparent().remove(el)

def insert_after(p, text):
    new_p = OxmlElement("w:p")
    p._p.addnext(new_p)
    q = Paragraph(new_p, p._parent)
    q.style = p.style
    q.add_run(text)
    return q

# Experiment 1: preserve the existing trial values, but make the interaction clearer.
p = find_para("click", "parallel paths", "fixed settings")
append_once(p, " The values listed under Fixed Parameters are already set in the simulation and do not need to be entered.", "Fixed Parameters")

p = find_para("for trial 1", "0.00", "8 w")
append_once(p, " Under Adjustable Parameters, click inside each input box, type the trial value shown above, and then click “Run and record trial.” The allowed range is printed directly below each box.", "Adjustable Parameters")

p = find_para("for trial 2", "0.85", "0 w")
append_once(p, " Change only the two adjustable input boxes to the Trial 2 values, then click “Run and record trial” and copy the displayed outputs into Table 1.", "adjustable input boxes")

p = find_para("for trial 3", "0.85", "8 w")
append_once(p, " Change the two adjustable input boxes to the Trial 3 values, click “Run and record trial,” and copy the displayed outputs into Table 1.", "adjustable input boxes")

# Experiment 2: make the thickness-entry workflow explicit without changing the experiment.
p = find_para("continue to series conduction", "fixed settings")
append_once(p, " These values are shown under Fixed Parameters and remain unchanged for Trials 4-6.", "Fixed Parameters")

p = find_para("run trial 4", "0.000")
append_once(p, " Click the Insulation thickness box under Adjustable Parameters, type 0.000, then click “Run and record trial.” The allowed range is shown below the box.", "Adjustable Parameters")

p = find_para("run trial 5", "0.100")
append_once(p, " Replace the value in the same Insulation thickness box with 0.100, click “Run and record trial,” and copy the displayed outputs into Table 2.", "Replace the value")

p = find_para("run trial 6", "0.150")
append_once(p, " Replace the value with 0.150, click “Run and record trial,” and copy the displayed outputs into Table 2.", "Replace the value")

# Experiment 3: replace the boundary-search workflow with two prescribed adjacent tests.
p = find_para("confirm the fixed settings", "800", "0.080")
append_once(p, " These values are shown under Fixed Parameters and should not be changed.", "Fixed Parameters")

replace_para(
    find_para("use insulation-thickness increments", "0.005"),
    "Under Adjustable Parameter, click the Insulation thickness box, enter 0.105 m, and click “Run and record trial.” Copy the displayed outputs to Trial 7 in Table 3 and classify the result by comparing T_s with the 60 °C safety limit."
)
replace_para(
    find_para("click", "record current trial", "trial 7"),
    "Change only the insulation thickness to 0.110 m. Click “Run and record trial,” copy the displayed outputs to Trial 8 in Table 3, and classify the result as PASS or FAIL."
)
replace_para(
    find_para("reduce", "0.005", "trial 8"),
    "Compare Trials 7 and 8. Under the fixed model conditions, one of these adjacent 0.005 m settings should fail the 60 °C requirement and the other should pass."
)
replace_para(
    find_para("mark each run", "pass or fail"),
    "Use the two recorded results to identify the minimum selectable safe insulation thickness for the 0.005 m thickness increments used in the simulation."
)

# Update Table 3 prescribed thicknesses.
table3 = None
for table in reversed(doc.tables):
    if len(table.rows) >= 3 and len(table.rows[0].cells) >= 7:
        table3 = table
        break
if table3 is None:
    raise RuntimeError("Safety Table 3 not found")

# The safety table has one header row followed by Trials 7 and 8.
table3.rows[1].cells[0].text = "7"
table3.rows[1].cells[1].text = "0.105"
table3.rows[2].cells[0].text = "8"
table3.rows[2].cells[1].text = "0.110"

# Align the post-experiment question with the two prescribed trials.
replace_para(
    find_para("why must the setting immediately below trial 7"),
    "Compare Trials 7 and 8. Which thickness passes and which fails? What do these two adjacent results show about the minimum selectable safe insulation thickness?"
)

# Replace the old Trial-8-is-one-step-lower equation block after the Experiment 3 safety statement.
safety_intro = find_para("use", "safety decision", "touchable outer-surface")
q9 = find_para("why is the outer-surface temperature", "touch-safety")
paras = doc.paragraphs
i0, i1 = paras.index(safety_intro), paras.index(q9)
between = paras[i0+1:i1]
mathish = [p for p in between if "oMath" in p._p.xml or not p.text.strip()]
if mathish:
    replace_para(mathish[0], "Safety comparison: T_s <= 60 °C; Trial 7: t_ins = 0.105 m; Trial 8: t_ins = 0.110 m.")
    for p in mathish[1:]:
        remove_para(p)
else:
    insert_after(safety_intro, "Safety comparison: T_s <= 60 °C; Trial 7: t_ins = 0.105 m; Trial 8: t_ins = 0.110 m.")

# Align Homework 3 with the new two-thickness experiment.
hw3 = find_para("question 3", "confirm the minimum selectable safe thickness")
replace_para(
    hw3,
    "Question 3. Confirm the minimum selectable safe thickness using the two prescribed safety trials and the requirement T_s <= 60 °C."
)
trial_report = find_para("for trials 7 and 8", "report the insulation thickness")
replace_para(
    trial_report,
    "For Trials 7 and 8, report the insulation thickness and outer-surface temperature and classify each trial as Pass or Fail. Confirm that 0.105 m fails and 0.110 m passes under the fixed model conditions. Explain briefly why these adjacent 0.005 m settings show that 0.110 m is the minimum selectable safe thickness."
)
paras = doc.paragraphs
i0, i1 = paras.index(hw3), paras.index(trial_report)
between = paras[i0+1:i1]
mathish = [p for p in between if "oMath" in p._p.xml or not p.text.strip()]
if mathish:
    replace_para(mathish[0], "Required comparison: T_s <= 60 °C; t_ins,Trial 7 = 0.105 m; t_ins,Trial 8 = 0.110 m.")
    for p in mathish[1:]:
        remove_para(p)
else:
    insert_after(hw3, "Required comparison: T_s <= 60 °C; t_ins,Trial 7 = 0.105 m; t_ins,Trial 8 = 0.110 m.")

# Save the updated worksheet in place for deployment.
doc.save(DOCX)

# Keep the legacy alias synchronized inside the deployed site.
legacy = BASE / "Panel_Heat_Transfer_Worksheet.docx"
legacy.write_bytes(DOCX.read_bytes())

print("Heat-transfer worksheet updated successfully.")
